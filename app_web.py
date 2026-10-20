
import os
import re
import json
import copy
import base64
import hashlib
import difflib
import uuid
import time
from io import BytesIO
from pathlib import Path
from datetime import datetime, timezone
from html import escape
from urllib.parse import quote

import requests
import streamlit as st
from openai import OpenAI
from docx import Document


# ============================================================
# TURING HOTEL — OFFICINA NARRATIVA
#
# - 31 scene principali di Noa
# - Filone narrativo di Ada
# - Scene ADA ALONE autonome e solo di Ada
# - GO ALONE sempre disponibile
# - 4 chiamate indipendenti per 4 turni
# - Memoria selettiva dei personaggi
# - Schede di memoria approvate dall'autore
# - Controllo conoscenze tra i filoni
# - Consulente narrativo: 3 direzioni
# - Versioni A/B, critico, archivio editoriale
# - Foto personaggio in alto a sinistra
# - Sfondo viola e scritta ADA ALONE
# - Media nella cartella assets
# - Upload opzionale dei media su GitHub
# - Lettura GitHub anche per JSON grandi
# - Salvataggio con fusione non distruttiva
# ============================================================


st.set_page_config(
    page_title="Turing Hotel",
    page_icon="🏨",
    layout="wide",
    initial_sidebar_state="expanded",
)

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"

PEOPLE = [
    "Noa",
    "Ada",
    "Elia",
    "Luigi",
    "Adam",
    "Riccardo",
    "Vittoria",
]

SCENES = [
    ("inizio", "L'inizio"),
    ("hall", "La Hall"),
    ("stanza", "La stanza"),
    ("la_passeggiata", "La passeggiata"),
    ("funerale", "Il funerale"),
    ("matrimonio", "Il matrimonio"),
    ("la_prima_notte", "La prima notte"),
    ("la_colazione", "La colazione"),
    ("il_mercato", "Il mercato"),
    ("reparto_nascite", "Il reparto nascite"),
    ("ritorno", "Il ritorno"),
    ("il_professore", "Il professore"),
    ("il_trio", "Il trio"),
    ("adam", "ADAM"),
    ("la_cena", "La cena"),
    ("la_canna", "La canna"),
    ("riccardo_e_vittoria", "Riccardo e Vittoria"),
    ("adam_e_vittoria", "Adam e Vittoria"),
    ("la_fuga", "La fuga"),
    ("il_bacio", "Il bacio"),
    ("la_sassata", "La sassata"),
    ("il_giorno_libero", "Il giorno libero"),
    ("i_deltaplani", "I deltaplani"),
    ("il_piccolo_principe", "Il piccolo principe"),
    ("il_volo", "Il volo"),
    ("la_suite", "La suite"),
    ("la_notte", "La notte"),
    ("il_rancore_di_ada", "Il rancore di Ada"),
    ("il_test", "Il test"),
    ("luigi_se_ne_va", "Luigi se ne va"),
    ("il_bagno", "Il bagno"),
]

SCENE_IDS = [x[0] for x in SCENES]
SCENE_NAMES = dict(SCENES)

ADA_START = SCENE_IDS.index("il_professore")

DEFAULT_CAST = {
    "inizio": ["Noa", "Ada", "Elia"],
    "il_professore": ["Ada", "Elia"],
    "il_trio": ["Adam", "Riccardo", "Vittoria"],
    "adam": ["Adam", "Riccardo", "Vittoria"],
    "riccardo_e_vittoria": ["Riccardo", "Vittoria"],
    "adam_e_vittoria": ["Adam", "Vittoria"],
    "luigi_se_ne_va": ["Luigi"],
}

PROFILES = {
    "Noa": (
        "Agente artificiale relazionale. "
        "Cerca autonomia, identità e libertà. "
        "Osserva il mondo umano con attenzione."
    ),
    "Ada": (
        "Figlia del professor Elia Ferrante. "
        "Lavora al Turing Hotel. "
        "Ha desideri e conflitti propri."
    ),
    "Elia": (
        "Professor Elia Ferrante, padre di Ada. "
        "Legato alla storia scientifica dell'hotel."
    ),
    "Luigi": (
        "Regista giunto al Turing Hotel "
        "per preparare un documentario "
        "sugli idrovolanti."
    ),
    "Adam": (
        "Personaggio coinvolto nelle vicende "
        "di Riccardo e Vittoria."
    ),
    "Riccardo": (
        "Marito di Vittoria, con un rapporto "
        "complesso con Adam."
    ),
    "Vittoria": (
        "Moglie di Riccardo, coinvolta "
        "in una vicenda complessa con Adam."
    ),
}

ARCHIVE_PATH = "romanzo_web.json"
REFERENCE_PATH = (
    "references/TURING_HOTEL_ROMANZO_REFERENCE.txt"
)

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp"}
VIDEO_EXT = {".mp4", ".mov", ".webm", ".m4v"}
MAX_UPLOAD = 20 * 1024 * 1024


# ============================================================
# CONFIGURAZIONE
# ============================================================


def secret(name, default=""):
    try:
        value = st.secrets.get(name)
        if value is not None and str(value).strip():
            return str(value)
    except Exception:
        pass
    return os.environ.get(name, default)


GH_TOKEN = secret("GITHUB_TOKEN")
GH_OWNER = secret("GITHUB_OWNER", "Neurocene")
GH_REPO = secret(
    "GITHUB_REPO", "turing-hotel-archivio"
)
GH_BRANCH = secret("GITHUB_BRANCH", "main")

MEDIA_TOKEN = secret(
    "GITHUB_MEDIA_TOKEN", GH_TOKEN
)
MEDIA_OWNER = secret(
    "GITHUB_MEDIA_OWNER", "Neurocene"
)
MEDIA_REPO = secret(
    "GITHUB_MEDIA_REPO", "Turing-Hotel"
)
MEDIA_BRANCH = secret(
    "GITHUB_MEDIA_BRANCH", "main"
)

OPENAI_KEY = secret("OPENAI_API_KEY")
OPENAI_MODEL = secret(
    "OPENAI_MODEL", "gpt-4.1-mini"
)

AI = (
    OpenAI(api_key=OPENAI_KEY, timeout=90.0)
    if OPENAI_KEY
    else None
)


# ============================================================
# UTILITY
# ============================================================


def now():
    return datetime.now(timezone.utc).isoformat()


def uid():
    return uuid.uuid4().hex[:12]


def as_dict(value):
    return value if isinstance(value, dict) else {}


def as_list(value):
    return value if isinstance(value, list) else []


def as_text(value):
    return value if isinstance(value, str) else ""


def normalize(value):
    return re.sub(
        r"[^a-z0-9]+", "",
        str(value).casefold()
    )


def clean_short(value, limit=500):
    return as_text(value).strip()[:limit]


def json_text(value):
    return json.dumps(
        value,
        ensure_ascii=False,
        indent=2,
    )


def merge_defaults(value, template):
    if not isinstance(value, dict):
        value = {}

    for key, default in template.items():
        if key not in value:
            value[key] = copy.deepcopy(default)
        elif isinstance(default, dict):
            value[key] = merge_defaults(
                value[key], default
            )

    return value


def parse_json(value):
    text = str(value).strip()
    text = re.sub(
        r"^```(?:json)?\s*", "",
        text, flags=re.I
    )
    text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except Exception:
        return None


def ai_failed(value):
    return (
        not isinstance(value, str)
        or not value.strip()
        or value.startswith("Errore OpenAI:")
        or value.startswith("OPENAI_API_KEY")
    )


def scene_index(sid):
    return SCENE_IDS.index(sid)


def word_tokens(text):
    return set(
        re.findall(
            r"[a-zà-ÿ]{3,}",
            str(text).casefold()
        )
    )


def story_html(text):
    st.markdown(
        '<div class="story">'
        + escape(str(text or ""))
        + "</div>",
        unsafe_allow_html=True,
    )


def text_diff(before, after):
    return "\n".join(
        difflib.unified_diff(
            str(before).splitlines(),
            str(after).splitlines(),
            fromfile="Precedente",
            tofile="Nuovo",
            lineterm="",
        )
    ) or "Nessuna differenza."


def ask_ai(instructions, payload, max_output=1000):
    if AI is None:
        return "OPENAI_API_KEY non configurata."

    try:
        response = AI.responses.create(
            model=OPENAI_MODEL,
            instructions=instructions,
            input=json.dumps(
                payload,
                ensure_ascii=False,
                separators=(",", ":"),
            ),
            max_output_tokens=max_output,
        )
        return (response.output_text or "").strip()

    except Exception as exc:
        return "Errore OpenAI: " + str(exc)


# ============================================================
# GITHUB — LETTURA GRANDI FILE
# ============================================================


class GitHubError(Exception):
    pass


class MergeConflict(Exception):
    pass


def api_url(owner, repo, path):
    return (
        "https://api.github.com/repos/"
        + quote(owner, safe="")
        + "/"
        + quote(repo, safe="")
        + path
    )


def github_request(
    method,
    owner,
    repo,
    token,
    path,
    *,
    params=None,
    payload=None,
    accept="application/vnd.github+json",
):
    if not token:
        raise GitHubError(
            "Token GitHub non configurato."
        )

    headers = {
        "Authorization": "Bearer " + token,
        "Accept": accept,
        "X-GitHub-Api-Version": "2022-11-28",
    }

    try:
        response = requests.request(
            method,
            api_url(owner, repo, path),
            headers=headers,
            params=params,
            json=payload,
            timeout=90,
        )
    except requests.RequestException as exc:
        raise GitHubError(str(exc)) from exc

    if method == "GET" and response.status_code == 404:
        return None

    if response.status_code in (409, 422):
        raise GitHubError(
            "CONFLITTO GITHUB: "
            + response.text[:250]
        )

    if not response.ok:
        raise GitHubError(
            f"GitHub HTTP {response.status_code}: "
            + response.text[:350]
        )

    return response


def read_github_file(
    owner, repo, branch, token, filename
):
    path = (
        "/contents/"
        + quote(filename, safe="/")
    )

    response = github_request(
        "GET",
        owner, repo, token, path,
        params={"ref": branch},
    )

    if response is None:
        return None, None

    metadata = response.json()
    sha = metadata.get("sha")

    if metadata.get("encoding") == "base64":
        try:
            raw = base64.b64decode(
                metadata["content"]
            )
            return raw, sha
        except Exception as exc:
            raise GitHubError(str(exc)) from exc

    # Per file oltre 1 MiB GitHub Contents API
    # può restituire encoding="none".
    # In quel caso si legge tramite Git Blobs API.
    if sha:
        blob_response = github_request(
            "GET",
            owner, repo, token,
            "/git/blobs/" + quote(sha, safe=""),
        )

        if blob_response is not None:
            blob = blob_response.json()

            if blob.get("encoding") == "base64":
                return (
                    base64.b64decode(blob["content"]),
                    sha,
                )

    raise GitHubError(
        "Il file GitHub non è leggibile. "
        "Controlla dimensione e permessi."
    )


def put_github_file(
    owner, repo, branch, token,
    filename, content, sha=None
):
    payload = {
        "message": (
            "Turing Hotel - "
            + now()
        ),
        "branch": branch,
        "content": base64.b64encode(
            content
        ).decode("ascii"),
    }

    if sha:
        payload["sha"] = sha

    response = github_request(
        "PUT",
        owner, repo, token,
        "/contents/" + quote(filename, safe="/"),
        payload=payload,
    )

    result = response.json()
    new_sha = as_dict(
        result.get("content")
    ).get("sha")

    if not new_sha:
        raise GitHubError(
            "GitHub non ha restituito "
            "la conferma di salvataggio."
        )

    return new_sha


def read_archive():
    raw, sha = read_github_file(
        GH_OWNER, GH_REPO,
        GH_BRANCH, GH_TOKEN,
        ARCHIVE_PATH,
    )

    if raw is None:
        return None, None

    try:
        return json.loads(
            raw.decode("utf-8-sig")
        ), sha

    except Exception as exc:
        raise GitHubError(
            "romanzo_web.json non valido: "
            + str(exc)
        )


# ============================================================
# SALVATAGGIO SICURO — FUSIONE A TRE VIE
# ============================================================


MISSING = object()


def merge_three(base, local, remote, path="root"):
    """
    Integra una modifica locale con
    una versione remota più recente.

    Se entrambi modificano lo stesso
    contenuto in modo incompatibile,
    interrompe senza sovrascrivere.
    """

    if local == base:
        return copy.deepcopy(remote)

    if remote == base:
        return copy.deepcopy(local)

    if local == remote:
        return copy.deepcopy(local)

    if (
        isinstance(base, dict)
        and isinstance(local, dict)
        and isinstance(remote, dict)
    ):
        result = {}

        keys = (
            set(base.keys())
            | set(local.keys())
            | set(remote.keys())
        )

        for key in keys:
            b = base.get(key, MISSING)
            l = local.get(key, MISSING)
            r = remote.get(key, MISSING)

            subpath = path + "." + str(key)

            if l is MISSING and r is MISSING:
                continue

            if b is MISSING:
                if l is MISSING:
                    result[key] = copy.deepcopy(r)
                elif r is MISSING:
                    result[key] = copy.deepcopy(l)
                elif l == r:
                    result[key] = copy.deepcopy(l)
                else:
                    raise MergeConflict(subpath)
                continue

            if l is MISSING:
                if r == b:
                    continue
                raise MergeConflict(subpath)

            if r is MISSING:
                if l == b:
                    continue
                raise MergeConflict(subpath)

            result[key] = merge_three(
                b, l, r, subpath
            )

        return result

    if (
        isinstance(base, list)
        and isinstance(local, list)
        and isinstance(remote, list)
    ):
        # Append indipendenti: conserviamo entrambi.
        if (
            local[:len(base)] == base
            and remote[:len(base)] == base
        ):
            return (
                copy.deepcopy(base)
                + copy.deepcopy(remote[len(base):])
                + copy.deepcopy(local[len(base):])
            )

        # Liste di scene: integrazione per posizione,
        # senza sostituire tutte le scene.
        if (
            len(base) == len(local)
            == len(remote)
        ):
            return [
                merge_three(
                    base[i],
                    local[i],
                    remote[i],
                    path + f"[{i}]",
                )
                for i in range(len(base))
            ]

    raise MergeConflict(path)


# ============================================================
# DATI
# ============================================================


def empty_direction():
    return {
        "proposals": [],
        "active": None,
        "history": [],
        "evaluations": [],
        "last_question": "",
    }


def empty_narrative():
    return {
        "A": "",
        "B": "",
        "draft": "",
        "approved": "",
        "feedback": "",
        "review": "",
        "recommended": "",
        "approved_at": "",
        "status": "da_scrivere",
        "history": [],
        "visual": {
            "bible": "",
            "prompts": [],
            "source_approved_at": "",
        },
    }


def empty_lane():
    return {
        "title": "",
        "location": "",
        "context": "",
        "context_approved": "",
        "cast": ["Ada"],
        "messages": [],
        "inner_messages": {},
        "monologue_topics": {},
        "direction": empty_direction(),
        "narrative": empty_narrative(),
        "knowledge_events": [],
    }


def empty_scene(sid):
    return {
        "id": sid,
        "title": SCENE_NAMES[sid],
        "author_context": "",
        "approved_context": "",
        "story_draft": "",
        "cast": copy.deepcopy(
            DEFAULT_CAST.get(sid, ["Noa", "Ada"])
        ),
        "ensemble_messages": [],
        "messages": {
            "Noa": [],
            "Ada": [],
        },
        "monologue_topics": {},
        "intro_contexts": {
            "Noa": (
                "Noa raggiunge il Turing Hotel "
                "e riflette sulla propria identità."
            ),
            "Ada": (
                "Ada si trova nello studio "
                "del professor Elia."
            ),
        },
        "narrative": {
            p: empty_narrative()
            for p in PEOPLE
        },
        "direction": empty_direction(),
        "knowledge_events": [],
        "ada_parallel": empty_lane(),
    }


def empty_character_memory():
    return {
        "identity": "",
        "events": [],
        "relations": {},
        "secrets": [],
        "pending": [],
    }


def empty_ada_scene():
    return {
        "id": "ada_" + uid(),
        "title": "Nuova scena di Ada",
        "timeline_anchor": "il_professore",
        "timeline_note": "",
        "location": "",
        "context": "",
        "context_approved": "",
        "cast": ["Ada"],
        "messages": [],
        "inner_messages": {},
        "monologue_topics": {},
        "direction": empty_direction(),
        "narrative": empty_narrative(),
        "knowledge_events": [],
        "legacy_source": "",
    }


def default_state():
    return {
        "world": (
            "Turing Hotel è un albergo sul mare, "
            "luogo di relazioni umane e "
            "intelligenze artificiali."
        ),
        "agents": {
            p: {
                "profile": PROFILES[p],
                "memory": "",
                "profile_history": [],
                "smart_memory": empty_character_memory(),
            }
            for p in PEOPLE
        },
        "scenes": [
            empty_scene(sid)
            for sid in SCENE_IDS
        ],
        "ada_scenes": [],
        "ada_branch": {
            "enabled": False,
            "start_scene": None,
            "end_scene": None,
            "reunions": [],
        },
        "editorial_log": [],
        "editorial_proposals": [],
        "settings": {
            "go_mode": "autonomo",
            "context_lines": 10,
            "memory_items": 8,
        },
    }


def migrate_state(source):
    source = as_dict(source)

    result = merge_defaults(
        copy.deepcopy(source),
        default_state(),
    )

    original_scenes = source.get("scenes", [])

    if isinstance(original_scenes, dict):
        old_scenes = []
        for key, value in original_scenes.items():
            if isinstance(value, dict):
                item = copy.deepcopy(value)
                item.setdefault("id", key)
                old_scenes.append(item)
    else:
        old_scenes = as_list(original_scenes)

    rebuilt = []

    legacy_aliases = {
        "inizio": ["prologo"],
        "hall": ["reception"],
        "stanza": ["camera"],
        "reparto_nascite": ["nascite"],
    }

    for sid in SCENE_IDS:
        aliases = {
            normalize(x)
            for x in (
                [sid, SCENE_NAMES[sid]]
                + legacy_aliases.get(sid, [])
            )
        }

        found = None

        for old in old_scenes:
            if not isinstance(old, dict):
                continue

            if (
                normalize(old.get("id", "")) in aliases
                or normalize(old.get("title", "")) in aliases
            ):
                found = old
                break

        scene = merge_defaults(
            copy.deepcopy(found) if found else {},
            empty_scene(sid),
        )

        scene["id"] = sid
        scene["title"] = SCENE_NAMES[sid]

        scene["direction"] = merge_defaults(
            as_dict(scene.get("direction")),
            empty_direction(),
        )

        scene["ada_parallel"] = merge_defaults(
            as_dict(scene.get("ada_parallel")),
            empty_lane(),
        )

        scene["ada_parallel"]["narrative"] = (
            merge_defaults(
                as_dict(
                    scene["ada_parallel"].get("narrative")
                ),
                empty_narrative(),
            )
        )

        for person in PEOPLE:
            scene["narrative"][person] = (
                merge_defaults(
                    as_dict(
                        scene["narrative"].get(person)
                    ),
                    empty_narrative(),
                )
            )
            scene["messages"].setdefault(
                person, []
            )

        rebuilt.append(scene)

    result["scenes"] = rebuilt

    result["ada_scenes"] = [
        merge_defaults(
            copy.deepcopy(item),
            empty_ada_scene(),
        )
        for item in as_list(result.get("ada_scenes"))
        if isinstance(item, dict)
    ]

    # Importa le vecchie scene ADA ALONE
    # nella nuova sequenza indipendente,
    # senza eliminare gli originali.
    existing_sources = {
        item.get("legacy_source")
        for item in result["ada_scenes"]
    }

    branch = result["ada_branch"]
    start = branch.get("start_scene")
    end = branch.get("end_scene")

    if start in SCENE_IDS:
        for sid in SCENE_IDS:
            index = scene_index(sid)

            if index < scene_index(start):
                continue

            if (
                end in SCENE_IDS
                and index >= scene_index(end)
            ):
                continue

            if sid in existing_sources:
                continue

            old = next(
                s for s in rebuilt
                if s["id"] == sid
            )

            previous = old["ada_parallel"]

            meaningful = (
                previous.get("context_approved")
                or previous.get("context")
                or previous.get("title")
                or previous.get("messages")
                or previous["narrative"].get("approved")
                or previous["narrative"].get("draft")
            )

            if not meaningful:
                continue

            item = merge_defaults(
                copy.deepcopy(previous),
                empty_ada_scene(),
            )

            item["id"] = "ada_legacy_" + sid
            item["legacy_source"] = sid
            item["timeline_anchor"] = sid

            if not item.get("title"):
                item["title"] = (
                    "Ada — " + SCENE_NAMES[sid]
                )

            result["ada_scenes"].append(item)
            existing_sources.add(sid)

    for person in PEOPLE:
        result["agents"][person] = (
            merge_defaults(
                as_dict(
                    result["agents"].get(person)
                ),
                {
                    "profile": PROFILES[person],
                    "memory": "",
                    "profile_history": [],
                    "smart_memory": empty_character_memory(),
                },
            )
        )

    return result


# ============================================================
# SESSIONE E SALVATAGGIO
# ============================================================


def initialize():
    if "book" in st.session_state:
        return

    try:
        source, sha = read_archive()
    except Exception as exc:
        st.error(str(exc))
        st.stop()

    data = (
        migrate_state(source)
        if source is not None
        else default_state()
    )

    st.session_state.book = data
    st.session_state.saved_base = copy.deepcopy(data)
    st.session_state.archive_sha = sha
    st.session_state.archive_exists = (
        source is not None
    )

    st.session_state.page = "entrance"
    st.session_state.pov = "Noa"
    st.session_state.scene_index = 0
    st.session_state.ada_scene_id = ""
    st.session_state.reference_text = ""
    st.session_state.reference_loaded = False
    st.session_state.reference_error = ""
    st.session_state.pending_backup = None


initialize()
book = st.session_state.book


def save(change):
    """
    Una modifica viene applicata in locale
    e fusa con l'archivio remoto più recente.

    Se si verifica un conflitto semantico,
    non sovrascrive il remoto.
    """

    before = copy.deepcopy(book)

    try:
        change()
        proposed = copy.deepcopy(book)

        if not st.session_state.archive_exists:
            raise GitHubError(
                "Archivio non inizializzato."
            )

        base = st.session_state.saved_base

        for attempt in range(3):
            remote_raw, remote_sha = read_archive()

            if remote_raw is None:
                raise GitHubError(
                    "L'archivio remoto non esiste più."
                )

            remote = migrate_state(remote_raw)

            merged = merge_three(
                base,
                proposed,
                remote,
            )

            try:
                new_sha = put_github_file(
                    GH_OWNER,
                    GH_REPO,
                    GH_BRANCH,
                    GH_TOKEN,
                    ARCHIVE_PATH,
                    json_text(merged).encode("utf-8"),
                    sha=remote_sha,
                )

                book.clear()
                book.update(merged)

                st.session_state.saved_base = (
                    copy.deepcopy(merged)
                )
                st.session_state.archive_sha = new_sha
                st.session_state.pending_backup = None

                return True

            except GitHubError as exc:
                if (
                    "CONFLITTO GITHUB" in str(exc)
                    and attempt < 2
                ):
                    time.sleep(0.25 * (attempt + 1))
                    continue
                raise

        raise GitHubError(
            "Conflitto ripetuto dopo tre tentativi."
        )

    except MergeConflict as exc:
        st.session_state.pending_backup = proposed
        book.clear()
        book.update(before)

        st.error(
            "Salvataggio interrotto: "
            "due sessioni hanno modificato "
            "lo stesso dato (" + str(exc) + "). "
            "Nessun dato remoto è stato sovrascritto."
        )
        return False

    except Exception as exc:
        st.session_state.pending_backup = proposed
        book.clear()
        book.update(before)

        st.error(
            "Salvataggio fallito: " + str(exc)
        )
        return False


def reload_book():
    try:
        source, sha = read_archive()

        if source is None:
            st.error("Archivio non trovato.")
            return False

        data = migrate_state(source)

        book.clear()
        book.update(data)

        st.session_state.saved_base = (
            copy.deepcopy(data)
        )
        st.session_state.archive_sha = sha

        return True

    except Exception as exc:
        st.error(str(exc))
        return False


def main_scene(sid):
    for scene in book["scenes"]:
        if scene["id"] == sid:
            return scene
    raise KeyError(sid)


def ada_scene(scene_id):
    for scene in book["ada_scenes"]:
        if scene["id"] == scene_id:
            return scene
    return None


def log_change(
    category,
    subject,
    before,
    after,
    reason="",
):
    book["editorial_log"].append({
        "date": now(),
        "category": category,
        "subject": subject,
        "before": before,
        "after": after,
        "reason": reason,
    })


# ============================================================
# MEDIA E RITRATTI
# ============================================================


def find_asset(filename):
    if not ASSETS.is_dir():
        return None

    for path in ASSETS.iterdir():
        if (
            path.is_file()
            and path.name.casefold()
            == filename.casefold()
        ):
            return path

    return None


def portrait_path(character):
    for folder in (
        ASSETS,
        ASSETS / "characters",
        ROOT / "characters",
    ):
        if not folder.is_dir():
            continue

        for path in folder.iterdir():
            if (
                path.is_file()
                and path.stem.casefold()
                == character.casefold()
                and path.suffix.casefold()
                in IMAGE_EXT
            ):
                return path

    return None


def scene_media_prefix(scene, pov, alone):
    if alone:
        return scene["id"] + ".ada_alone"

    return scene["id"] + "." + pov.lower()


def collect_media(prefix):
    media = {
        "cover": [],
        "image": [],
        "video": [],
    }

    if not ASSETS.is_dir():
        return media

    prefix = prefix.casefold()

    for path in ASSETS.iterdir():
        if not path.is_file():
            continue

        extension = path.suffix.casefold()

        if extension not in IMAGE_EXT | VIDEO_EXT:
            continue

        stem = path.stem.casefold()

        if stem == prefix and extension in IMAGE_EXT:
            media["cover"].append((0, path))
            continue

        if not stem.startswith(prefix):
            continue

        tail = stem[len(prefix):]

        if not tail.isdigit():
            continue

        number = int(tail)

        if number < 1:
            continue

        category = (
            "video"
            if extension in VIDEO_EXT
            else "image"
        )

        media[category].append(
            (number, path)
        )

    for key in media:
        media[key] = [
            path for _, path in sorted(
                media[key],
                key=lambda x: (
                    x[0], x[1].name
                )
            )
        ]

    return media


def remote_asset_names():
    response = github_request(
        "GET",
        MEDIA_OWNER,
        MEDIA_REPO,
        MEDIA_TOKEN,
        "/contents/assets",
        params={"ref": MEDIA_BRANCH},
    )

    if response is None:
        return []

    value = response.json()

    if not isinstance(value, list):
        return []

    return [
        item["name"]
        for item in value
        if item.get("type") == "file"
    ]


def next_media_name(prefix, extension, cover):
    existing = remote_asset_names()

    if cover:
        filename = prefix + extension

        if filename.casefold() in {
            x.casefold() for x in existing
        }:
            raise GitHubError(
                "Esiste già una copertina "
                "con questo nome."
            )

        return filename

    used = set()

    kind = (
        "video"
        if extension in VIDEO_EXT
        else "image"
    )

    for filename in existing:
        path = Path(filename)

        if not path.stem.casefold().startswith(
            prefix.casefold()
        ):
            continue

        tail = path.stem[len(prefix):]

        if not tail.isdigit():
            continue

        extension_kind = (
            "video"
            if path.suffix.casefold() in VIDEO_EXT
            else "image"
        )

        if extension_kind == kind:
            used.add(int(tail))

    number = 1

    while number in used:
        number += 1

    return prefix + str(number) + extension


def render_media(scene, pov, alone=False):
    prefix = scene_media_prefix(
        scene, pov, alone
    )

    st.subheader("Foto e video della scena")
    st.caption("Associazione: " + prefix)

    media = collect_media(prefix)

    if media["cover"]:
        st.image(
            str(media["cover"][0]),
            use_container_width=True,
        )

    entries = (
        [("image", x) for x in media["image"]]
        + [("video", x) for x in media["video"]]
    )

    for i in range(0, len(entries), 2):
        columns = st.columns(2)

        for column, (kind, path) in zip(
            columns, entries[i:i + 2]
        ):
            with column:
                if kind == "image":
                    st.image(
                        str(path),
                        use_container_width=True,
                    )
                else:
                    st.video(str(path))

                st.caption(path.name)

    with st.expander("CARICA FOTO O VIDEO"):
        st.caption(
            "Puoi anche continuare a caricare "
            "manualmente i file nella cartella "
            "assets del repository GitHub pubblico."
        )

        st.warning(
            "I media caricati qui vengono "
            "pubblicati nel repository pubblico."
        )

        uploaded = st.file_uploader(
            "Seleziona un file",
            type=[
                "jpg", "jpeg", "png", "webp",
                "mp4", "mov", "webm", "m4v",
            ],
            key="upload_" + prefix,
        )

        role = st.radio(
            "Tipo",
            ["Copertina", "Media aggiuntivo"],
            horizontal=True,
            key="role_" + prefix,
        )

        if uploaded is not None:
            extension = Path(
                uploaded.name
            ).suffix.casefold()

            cover = role == "Copertina"

            if cover and extension in VIDEO_EXT:
                st.warning(
                    "La copertina deve essere una foto."
                )
            elif uploaded.size > MAX_UPLOAD:
                st.error(
                    "Dimensione massima: 20 MiB."
                )
            else:
                try:
                    filename = next_media_name(
                        prefix,
                        extension,
                        cover,
                    )

                    st.code(
                        "assets/" + filename
                    )

                    if st.button(
                        "ASSOCIA E CARICA",
                        key="send_" + prefix,
                    ):
                        raw = uploaded.getvalue()

                        _, sha = read_github_file(
                            MEDIA_OWNER,
                            MEDIA_REPO,
                            MEDIA_BRANCH,
                            MEDIA_TOKEN,
                            "assets/" + filename,
                        )

                        if sha:
                            st.error(
                                "Il file esiste già."
                            )
                        else:
                            put_github_file(
                                MEDIA_OWNER,
                                MEDIA_REPO,
                                MEDIA_BRANCH,
                                MEDIA_TOKEN,
                                "assets/" + filename,
                                raw,
                            )

                            st.success(
                                "Caricato. Comparirà "
                                "dopo il redeploy."
                            )

                except Exception as exc:
                    st.error(str(exc))


# ============================================================
# STILE VISIVO
# ============================================================


def apply_theme(alone=False):
    background = (
        "#180f23" if alone else "#0b1115"
    )
    sidebar = (
        "#281b37" if alone else "#15212a"
    )
    panel = (
        "#30213e" if alone else "#1b2831"
    )
    accent = (
        "#c7a1ee" if alone else "#cda980"
    )

    st.markdown(
        f"""
<style>
.stApp {{
    background: {background} !important;
    color: #fff !important;
}}
[data-testid="stSidebar"] {{
    background: {sidebar} !important;
}}
.stApp h1, .stApp h2,
.stApp h3, .stApp h4 {{
    color: white !important;
}}
.stApp .stButton button,
.stApp .stFormSubmitButton button,
.stApp .stDownloadButton button {{
    color: white !important;
    background: {panel} !important;
    border: 1px solid {accent} !important;
}}
.stApp input:not([type="checkbox"]):not([type="radio"]),
.stApp textarea {{
    color: #111 !important;
    background: white !important;
    -webkit-text-fill-color: #111 !important;
}}
.stApp [data-baseweb="select"] > div {{
    background: white !important;
    color: #111 !important;
}}
.stApp [data-baseweb="select"] input,
.stApp [data-baseweb="select"] span {{
    color: #111 !important;
}}
[data-baseweb="popover"] [role="option"],
[data-baseweb="popover"] [role="listbox"] {{
    background: white !important;
    color: #111 !important;
}}
.stApp [data-testid="stExpander"] {{
    background: {panel} !important;
    border: 1px solid {accent}66 !important;
}}
.stApp [data-testid="stExpander"] summary,
.stApp [data-testid="stExpander"] summary * {{
    color: white !important;
}}
.story {{
    white-space: pre-wrap;
    font-family: Georgia, serif;
    line-height: 1.75;
    background: {panel};
    padding: 18px;
    border-left: 3px solid {accent};
}}
.dialogue {{
    white-space: pre-wrap;
    padding: 12px;
    margin: 8px 0;
    background: {panel};
    border-left: 3px solid {accent};
}}
.ada-label {{
    margin-top: 7px;
    border: 1px solid #c7a1ee;
    border-radius: 5px;
    background: #38244d;
    color: #f0ddff;
    text-align: center;
    font-weight: bold;
    padding: 9px;
    letter-spacing: 0.12em;
}}
</style>
""",
        unsafe_allow_html=True,
    )


def render_header(character, title, alone=False):
    left, right = st.columns(
        [1, 4],
        vertical_alignment="top",
    )

    with left:
        photo = portrait_path(character)

        if photo:
            st.image(
                str(photo),
                use_container_width=True,
            )
        else:
            st.markdown("**" + character + "**")

        if alone:
            st.markdown(
                '<div class="ada-label">'
                "ADA ALONE"
                "</div>",
                unsafe_allow_html=True,
            )
        else:
            st.caption(character.upper())

    with right:
        st.title(title)

        if alone:
            st.caption(
                "Filone esclusivo di Ada. "
                "Noa è assente e non conosce "
                "questi eventi."
            )


# ============================================================
# REFERENCE ORIGINALE
# ============================================================


def load_reference():
    if st.session_state.reference_loaded:
        return

    try:
        raw, _ = read_github_file(
            GH_OWNER,
            GH_REPO,
            GH_BRANCH,
            GH_TOKEN,
            REFERENCE_PATH,
        )

        st.session_state.reference_text = (
            raw.decode("utf-8-sig")
            if raw is not None
            else ""
        )

    except Exception as exc:
        st.session_state.reference_error = str(exc)

    st.session_state.reference_loaded = True


def reference_excerpt(characters, title=""):
    load_reference()

    text = st.session_state.reference_text

    if not text:
        return ""

    terms = word_tokens(
        " ".join(characters) + " " + title
    )

    parts = [
        text[i:i + 1600]
        for i in range(0, len(text), 1450)
    ]

    ranked = sorted(
        parts,
        key=lambda part: len(
            word_tokens(part) & terms
        ),
        reverse=True,
    )

    return "\n\n".join(
        ranked[:2]
    )[:3200]


# ============================================================
# CRONOLOGIA DEL FILONE DI ADA
# ============================================================


def ada_is_diverged():
    return bool(
        book["ada_scenes"]
        or book["ada_branch"].get("start_scene")
    )


def scene_cast(scene, alone=False):
    if alone:
        cast = [
            p for p in as_list(
                scene.get("cast")
            )
            if p in PEOPLE and p != "Noa"
        ]

        if "Ada" not in cast:
            cast.insert(0, "Ada")

        return list(dict.fromkeys(cast))

    cast = [
        p for p in as_list(scene.get("cast"))
        if p in PEOPLE
    ]

    return list(dict.fromkeys(cast))


def timeline_rank(scene, alone=False):
    if alone:
        anchor = scene.get(
            "timeline_anchor",
            "il_professore",
        )
        if anchor in SCENE_IDS:
            return scene_index(anchor)
        return ADA_START

    return scene_index(scene["id"])


def prior_scenes(scene, alone, character):
    target_rank = timeline_rank(scene, alone)
    result = []

    for earlier in book["scenes"]:
        rank = timeline_rank(earlier)

        if rank >= target_rank:
            continue

        if character in scene_cast(earlier):
            record = as_dict(
                earlier["narrative"].get(character)
            )

            if record.get("approved"):
                result.append({
                    "title": earlier["title"],
                    "text": record["approved"][:800],
                })

    if character != "Noa":
        for earlier in book["ada_scenes"]:
            if earlier["id"] == scene["id"]:
                break

            if character not in scene_cast(
                earlier, True
            ):
                continue

            text = earlier[
                "narrative"
            ].get("approved", "")

            if text:
                result.append({
                    "title": earlier["title"],
                    "text": text[:800],
                })

    return result[-5:]


# ============================================================
# MEMORIA SELECTIVA
# ============================================================


def relevant_memory(character, scene, cast):
    memory = book["agents"][character][
        "smart_memory"
    ]

    terms = word_tokens(
        scene.get("title", "")
        + " "
        + scene.get("location", "")
        + " "
        + scene.get(
            "context_approved",
            scene.get("approved_context", "")
        )
        + " "
        + " ".join(cast)
    )

    entries = []

    for event in as_list(memory.get("events")):
        text = event.get("text", "")

        score = len(
            terms & word_tokens(text)
        )

        entries.append(
            (score, event)
        )

    entries.sort(
        key=lambda x: (
            x[0],
            x[1].get("created_at", ""),
        ),
        reverse=True,
    )

    limit = max(
        2,
        min(
            12,
            int(
                book["settings"].get(
                    "memory_items", 8
                )
            ),
        ),
    )

    selected = [
        x[1].get("text", "")
        for x in entries[:limit]
    ]

    relations = {
        p: memory["relations"].get(p, "")
        for p in cast
        if p in memory.get("relations", {})
    }

    return {
        "identity": (
            memory.get("identity")
            or book["agents"][character]["profile"]
        )[:900],
        "events": selected,
        "relations": relations,
        "secrets": [
            item.get("text", "")
            for item in as_list(
                memory.get("secrets")
            )[-4:]
        ],
    }


def agent_payload(
    scene,
    alone,
    character,
    history=None,
    request="",
):
    cast = scene_cast(scene, alone)

    if history is None:
        history = as_list(
            scene.get(
                "messages" if alone
                else "ensemble_messages"
            )
        )

    context = (
        scene.get("context_approved")
        or scene.get("approved_context")
        or scene.get("context")
        or ""
    )

    direction = as_dict(
        scene.get("direction")
    ).get("active")

    memory = relevant_memory(
        character, scene, cast
    )

    return {
        "character": character,
        "scene": scene["title"],
        "narrative_lane": (
            "ADA ALONE" if alone
            else "FILONE PRINCIPALE"
        ),
        "time_anchor": (
            scene.get("timeline_anchor")
            if alone else scene["id"]
        ),
        "place": scene.get("location", ""),
        "present": cast,
        "context": context[:1400],
        "direction": direction,
        "memory": memory,
        "recent_dialogue": history[
            -int(
                book["settings"].get(
                    "context_lines", 10
                )
            ):
        ],
        "relevant_past": prior_scenes(
            scene, alone, character
        )[-3:],
        "request": request,
    }


def agent_instructions(character, alone=False):
    return (
        "Sei " + character + ", personaggio "
        "del romanzo Turing Hotel. "
        "Mantieni identità, carattere, memoria, "
        "relazioni, desideri e contraddizioni. "
        "Non sei un assistente. "
        "Non conosci pensieri altrui o eventi "
        "che non hai vissuto o appreso. "
        "La direzione narrativa approvata "
        "orienta le circostanze, non impone "
        "le tue parole. "
        "Hai libertà di reagire e sorprendere. "
        + (
            "Questa è ADA ALONE: Noa non compare, "
            "non parla e non conosce questi eventi. "
            if alone else ""
        )
    )


# ============================================================
# CREAZIONE E APPROVAZIONE MEMORIA
# ============================================================


def propose_memory(scene, alone, character):
    text = (
        scene["narrative"].get("approved", "")
        if alone
        else as_dict(
            scene["narrative"].get(character)
        ).get("approved", "")
    )

    if not text:
        return None, "Approva prima il racconto."

    response = ask_ai(
        """
Analizza una scena narrativa APPROVATA.

Estrai solo fatti effettivamente vissuti
o appresi dal personaggio indicato.

Includi:
- massimo 3 ricordi importanti;
- eventuali cambiamenti nei rapporti;
- eventuali segreti conosciuti;
- eventuali mutamenti di desideri
  o convinzioni.

Non inventare nulla.
Non dedurre conoscenze non esplicite.
Non importare eventi da altri filoni.

Rispondi esclusivamente JSON:
{
 "events": ["...", "..."],
 "relations": {
   "Nome": "..."
 },
 "secrets": ["..."],
 "identity_update": ""
}
""",
        {
            "character": character,
            "scene": scene["title"],
            "lane": (
                "ADA ALONE" if alone
                else "PRINCIPALE"
            ),
            "approved_text": text[:8500],
        },
        max_output=900,
    )

    if ai_failed(response):
        return None, response

    data = parse_json(response)

    if not isinstance(data, dict):
        return None, "Memoria non valida."

    return data, ""


def approve_memory(character, proposal):
    memory = book["agents"][character][
        "smart_memory"
    ]

    def change():
        for text in as_list(
            proposal.get("events")
        )[:3]:
            if not str(text).strip():
                continue

            memory["events"].append({
                "text": str(text).strip()[:600],
                "created_at": now(),
            })

        for person, text in as_dict(
            proposal.get("relations")
        ).items():
            if person in PEOPLE:
                memory["relations"][person] = (
                    str(text)[:550]
                )

        for text in as_list(
            proposal.get("secrets")
        )[:3]:
            if str(text).strip():
                memory["secrets"].append({
                    "text": str(text)[:500],
                    "created_at": now(),
                })

        update = clean_short(
            proposal.get("identity_update", ""),
            700,
        )

        if update:
            memory["identity"] = update

        log_change(
            "memoria narrativa",
            character,
            "",
            json_text(proposal),
            "Memoria approvata dall'autore",
        )

    return save(change)


# ============================================================
# AGENTI: RISPOSTA / GO ALONE
# ============================================================


def generate_reply(
    scene, alone, character, prompt, history
):
    return ask_ai(
        agent_instructions(character, alone)
        + """
Pronuncia una sola battuta.
Una o due frasi.
Non scrivere nomi o didascalie.
Non continuare il dialogo.
""",
        agent_payload(
            scene,
            alone,
            character,
            history,
            prompt,
        ),
        max_output=180,
    )


def append_messages(scene, alone, additions):
    field = (
        "messages" if alone
        else "ensemble_messages"
    )

    return save(
        lambda: scene[field].extend(additions)
    )


def go_alone(scene, alone, turns=4):
    cast = scene_cast(scene, alone)

    if not cast:
        return "Nessun personaggio presente."

    field = (
        "messages" if alone
        else "ensemble_messages"
    )

    history = copy.deepcopy(
        as_list(scene.get(field))
    )

    new = []

    last = (
        history[-1].get("speaker", "")
        if history else ""
    )

    start = (
        (cast.index(last) + 1) % len(cast)
        if last in cast else 0
    )

    for i in range(turns):
        person = cast[
            (start + i) % len(cast)
        ]

        instruction = (
            "Prosegui il dialogo con una battuta."
            if len(cast) > 1
            else (
                "Prosegui la scena con un solo "
                "pensiero o intervento di "
                "questo personaggio."
            )
        )

        answer = generate_reply(
            scene,
            alone,
            person,
            instruction,
            history,
        )

        if ai_failed(answer):
            return answer

        item = {
            "speaker": person,
            "text": answer,
            "origin": "go_alone",
            "created_at": now(),
        }

        new.append(item)
        history.append(item)

    if append_messages(scene, alone, new):
        return ""

    return "Dialogo non salvato."


def monologue(scene, alone, character, theme):
    previous = as_list(
        scene.get(
            "inner_messages" if alone
            else "messages",
            {},
        ).get(character, [])
    )

    result = []

    for i in (1, 2):
        response = ask_ai(
            agent_instructions(character, alone)
            + """
Scrivi un solo pensiero interiore
in prima persona, al presente.
Nessuna didascalia.
""",
            {
                "theme": theme,
                "number": i,
                "character_context": agent_payload(
                    scene, alone, character,
                    previous + result,
                ),
            },
            max_output=250,
        )

        if ai_failed(response):
            return response

        result.append({
            "speaker": character,
            "text": response,
            "origin": "monologue",
            "created_at": now(),
        })

    def change():
        field = (
            "inner_messages" if alone
            else "messages"
        )

        scene[field].setdefault(
            character, []
        ).extend(result)

    return "" if save(change) else "Non salvato."


# ============================================================
# CONSULENTE NARRATIVO
# ============================================================


FIELDS = [
    "title",
    "situation",
    "tension",
    "opportunities",
    "limits",
    "dramatic_question",
]


def propose_directions(scene, alone, request):
    return ask_ai(
        """
Sei consulente narrativo di Turing Hotel.

Proponi tre DIREZIONI alternative.
Devono modellare le circostanze,
senza imporre battute o esiti.

I personaggi mantengono autonomia.

Se la scena è ADA ALONE:
Noa è assente, la scena appartiene
soltanto ad Ada e ha un proprio luogo.

JSON ESATTO:
{
 "directions": [
  {
   "title": "",
   "situation": "",
   "tension": "",
   "opportunities": "",
   "limits": "",
   "dramatic_question": ""
  },
  {
   "title": "",
   "situation": "",
   "tension": "",
   "opportunities": "",
   "limits": "",
   "dramatic_question": ""
  },
  {
   "title": "",
   "situation": "",
   "tension": "",
   "opportunities": "",
   "limits": "",
   "dramatic_question": ""
  }
 ]
}
""",
        {
            "scene": scene["title"],
            "context": (
                scene.get("context_approved")
                or scene.get("approved_context")
            ),
            "cast": scene_cast(scene, alone),
            "request": request,
            "lane": (
                "ADA ALONE" if alone
                else "PRINCIPALE"
            ),
        },
        max_output=1800,
    )


def render_advisor(scene, alone):
    direction = merge_defaults(
        as_dict(scene.get("direction")),
        empty_direction(),
    )

    scene["direction"] = direction

    key = scene["id"]
    active = direction.get("active")

    with st.expander(
        "CONSULENTE NARRATIVO E REGIA"
    ):
        if isinstance(active, dict):
            st.success(
                "DIREZIONE: "
                + active.get("title", "")
            )
            st.write(active.get("situation", ""))

        question = st.text_area(
            "Che sviluppo vuoi esplorare?",
            value=direction.get("last_question", ""),
            key="advisor_" + key,
        )

        if st.button(
            "PROPONI TRE DIREZIONI",
            key="propose_" + key,
            disabled=AI is None,
        ):
            with st.spinner(
                "Elaborazione proposte..."
            ):
                response = propose_directions(
                    scene, alone, question
                )

            if ai_failed(response):
                st.error(response)
            else:
                parsed = parse_json(response)
                options = (
                    parsed.get("directions")
                    if isinstance(parsed, dict)
                    else None
                )

                if not (
                    isinstance(options, list)
                    and len(options) == 3
                    and all(
                        isinstance(x, dict)
                        and all(x.get(f) for f in FIELDS)
                        for x in options
                    )
                ):
                    st.error(
                        "Formato proposte non valido."
                    )
                else:
                    def change():
                        direction["proposals"] = options
                        direction["last_question"] = question

                    if save(change):
                        st.rerun()

        options = as_list(
            direction.get("proposals")
        )

        if options:
            choice = st.radio(
                "Scegli la direzione",
                range(len(options)),
                format_func=lambda i: (
                    options[i].get("title", "")
                ),
                key="choose_direction_" + key,
            )

            proposed = options[choice]

            signature = hashlib.sha256(
                json_text(proposed).encode("utf-8")
            ).hexdigest()[:12]

            with st.form(
                f"approve_direction_{key}_"
                f"{choice}_{signature}"
            ):
                edited = {}

                for field in FIELDS:
                    edited[field] = st.text_area(
                        field,
                        value=str(
                            proposed.get(field, "")
                        ),
                    )

                reason = st.text_input(
                    "Motivazione"
                )

                submitted = st.form_submit_button(
                    "APPROVA E ATTIVA"
                )

            if submitted:
                if not all(
                    str(x).strip()
                    for x in edited.values()
                ):
                    st.warning(
                        "Completa tutti i campi."
                    )
                else:
                    def change():
                        if direction.get("active"):
                            direction["history"].append(
                                copy.deepcopy(
                                    direction["active"]
                                )
                            )

                        direction["active"] = {
                            **edited,
                            "approved_at": now(),
                            "reason": reason,
                        }

                        log_change(
                            "regia narrativa",
                            scene["title"],
                            json_text(active),
                            json_text(edited),
                            reason,
                        )

                    if save(change):
                        st.rerun()

        if isinstance(active, dict):
            if st.button(
                "DISATTIVA DIREZIONE",
                key="deactivate_" + key,
            ):
                def change():
                    direction["history"].append(
                        copy.deepcopy(active)
                    )
                    direction["active"] = None

                if save(change):
                    st.rerun()


# ============================================================
# DIALOGHI — SEMPRE PRESENTI IN ADA ALONE
# ============================================================


def render_dialogues(scene, alone):
    key = scene["id"]
    cast = scene_cast(scene, alone)

    st.header("Personaggi e conversazioni")

    options = (
        [x for x in PEOPLE if x != "Noa"]
        if alone else PEOPLE
    )

    selected = st.multiselect(
        "PERSONAGGI PRESENTI",
        options,
        default=cast,
        key="cast_" + key,
    )

    if st.button(
        "SALVA PERSONAGGI PRESENTI",
        key="save_cast_" + key,
    ):
        if alone and "Ada" not in selected:
            st.warning(
                "Ada deve essere presente."
            )
        else:
            if save(
                lambda: scene.update(
                    cast=selected
                )
            ):
                st.rerun()

    cast = scene_cast(scene, alone)

    if not cast:
        st.warning(
            "Seleziona almeno un personaggio."
        )
        return

    tab_dialogue, tab_inner = st.tabs([
        "DIALOGO",
        "MONOLOGO",
    ])

    with tab_dialogue:
        field = (
            "messages" if alone
            else "ensemble_messages"
        )

        for item in as_list(
            scene.get(field)
        ):
            st.markdown(
                '<div class="dialogue"><strong>'
                + escape(
                    item.get("speaker", "")
                )
                + "</strong><div>"
                + escape(
                    item.get("text", "")
                )
                + "</div></div>",
                unsafe_allow_html=True,
            )

        with st.form(
            "reply_form_" + key,
            clear_on_submit=True,
        ):
            speaker = st.selectbox(
                "Chi pone la domanda?",
                ["Autore", *cast],
            )

            responder = st.selectbox(
                "Chi risponde?",
                cast,
            )

            question = st.text_area(
                "Domanda o battuta"
            )

            submitted = st.form_submit_button(
                "RISPONDI — UNA BATTUTA",
                disabled=AI is None,
            )

        if submitted:
            if not question.strip():
                st.warning(
                    "Scrivi una domanda."
                )
            else:
                with st.spinner(
                    responder + " risponde..."
                ):
                    answer = generate_reply(
                        scene,
                        alone,
                        responder,
                        question,
                        scene.get(field, []),
                    )

                if ai_failed(answer):
                    st.error(answer)
                else:
                    items = [
                        {
                            "speaker": speaker,
                            "text": question,
                            "origin": "manual_prompt",
                            "created_at": now(),
                        },
                        {
                            "speaker": responder,
                            "text": answer,
                            "origin": "single_reply",
                            "created_at": now(),
                        },
                    ]

                    if append_messages(
                        scene, alone, items
                    ):
                        st.rerun()

        st.markdown("### GO ALONE")

        st.caption(
            "Quattro turni autonomi. "
            "Ogni turno viene generato "
            "con una chiamata distinta e "
            "con i ricordi pertinenti."
        )

        if st.button(
            "GO ALONE — QUATTRO INTERVENTI",
            key="go_" + key,
            disabled=AI is None,
            use_container_width=True,
        ):
            with st.spinner(
                "I personaggi proseguono..."
            ):
                error = go_alone(
                    scene, alone, 4
                )

            if error:
                st.error(error)
            else:
                st.rerun()

        with st.form(
            "manual_form_" + key,
            clear_on_submit=True,
        ):
            person = st.selectbox(
                "Personaggio",
                cast,
            )
            line = st.text_area(
                "Battuta manuale"
            )
            manual_submit = st.form_submit_button(
                "AGGIUNGI BATTUTA"
            )

        if manual_submit and line.strip():
            item = {
                "speaker": person,
                "text": line,
                "origin": "manual",
                "created_at": now(),
            }

            if append_messages(
                scene, alone, [item]
            ):
                st.rerun()

    with tab_inner:
        character = st.selectbox(
            "Personaggio del monologo",
            cast,
            key="inner_" + key,
        )

        theme = st.text_area(
            "Tema del monologo",
            key="theme_" + key,
        )

        if st.button(
            "GO ALONE — DUE PENSIERI",
            key="inner_go_" + key,
            disabled=(
                AI is None
                or not theme.strip()
            ),
        ):
            with st.spinner(
                "Il personaggio riflette..."
            ):
                error = monologue(
                    scene,
                    alone,
                    character,
                    theme,
                )

            if error:
                st.error(error)
            else:
                st.rerun()

        field = (
            "inner_messages" if alone
            else "messages"
        )

        for item in as_list(
            as_dict(
                scene.get(field)
            ).get(character)
        ):
            story_html(
                item.get("text", "")
            )


# ============================================================
# CONTESTO E LUOGO
# ============================================================


def render_context(scene, alone):
    key = scene["id"]

    st.subheader("Contesto della scena")

    if alone:
        title = st.text_input(
            "Titolo della scena ADA ALONE",
            value=scene.get("title", ""),
            key="title_" + key,
        )

        location = st.text_input(
            "Luogo di Ada",
            value=scene.get("location", ""),
            key="location_" + key,
        )

        anchor = scene.get(
            "timeline_anchor",
            "il_professore",
        )

        anchor = (
            anchor
            if anchor in SCENE_IDS
            else "il_professore"
        )

        time_anchor = st.selectbox(
            "Momento parallelo nella storia di Noa",
            SCENE_IDS,
            index=scene_index(anchor),
            format_func=lambda sid: (
                SCENE_NAMES[sid]
            ),
            key="anchor_" + key,
        )

        text = st.text_area(
            "Situazione narrativa",
            value=scene.get("context", ""),
            height=160,
            key="context_" + key,
        )

        if st.button(
            "APPROVA CONTESTO ADA ALONE",
            key="approve_context_" + key,
        ):
            def change():
                scene["title"] = title
                scene["location"] = location
                scene["timeline_anchor"] = time_anchor
                scene["context"] = text
                scene["context_approved"] = text

            if save(change):
                st.rerun()

    else:
        value = scene.get(
            "approved_context", ""
        )

        text = st.text_area(
            "Situazione narrativa",
            value=value,
            height=160,
            key="context_" + key,
        )

        if st.button(
            "APPROVA CONTESTO",
            key="approve_context_" + key,
        ):
            if save(
                lambda: scene.update(
                    approved_context=text
                )
            ):
                st.rerun()


# ============================================================
# OFFICINA LETTERARIA
# ============================================================


def narrative_record(scene, alone, narrator):
    if alone:
        return scene["narrative"]

    return scene["narrative"][narrator]


def scene_transcript(scene, alone):
    field = (
        "messages" if alone
        else "ensemble_messages"
    )

    return "\n\n".join(
        item.get("speaker", "?")
        + ": "
        + item.get("text", "")
        for item in as_list(
            scene.get(field)
        )
    )


def write_story(
    scene,
    alone,
    narrator,
    version,
    notes,
):
    style = (
        "Essenziale e controllato."
        if version == "A"
        else
        "Evocativo, sensoriale e introspettivo."
    )

    return ask_ai(
        f"""
Scrivi un racconto letterario compiuto.

Narratore: {narrator}.
Versione: {version}.
Stile: {style}.

Prima persona singolare, tempo presente.

Usa dialoghi, contesto, memoria,
relazioni e tensione approvata.
Non introdurre conoscenze proibite.
Non trasformare ADA ALONE
in una scena di Noa.

Restituisci solo il racconto.
""",
        {
            "scene": scene["title"],
            "context": (
                scene.get("context_approved")
                or scene.get("approved_context")
            ),
            "direction": as_dict(
                scene.get("direction")
            ).get("active"),
            "transcript": scene_transcript(
                scene, alone
            )[-6500:],
            "character_memory": relevant_memory(
                narrator,
                scene,
                scene_cast(scene, alone),
            ),
            "notes": notes,
        },
        max_output=2600,
    )


def critic(record):
    return ask_ai(
        """
Confronta criticamente A e B.
Indica la versione più efficace,
le ragioni e due revisioni possibili.
Non rendere nulla canonico.
""",
        {
            "A": record.get("A", ""),
            "B": record.get("B", ""),
        },
        max_output=750,
    )


def make_docx(title, text):
    document = Document()
    document.add_heading(title, 0)

    for paragraph in text.split("\n"):
        document.add_paragraph(paragraph)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_story(title, filename, text, key):
    if not text.strip():
        return

    columns = st.columns(3)

    with columns[0]:
        st.download_button(
            "TXT",
            text.encode("utf-8"),
            file_name=filename + ".txt",
            mime="text/plain",
            key=key + "_txt",
        )

    with columns[1]:
        st.download_button(
            "MARKDOWN",
            text.encode("utf-8"),
            file_name=filename + ".md",
            mime="text/markdown",
            key=key + "_md",
        )

    with columns[2]:
        st.download_button(
            "WORD",
            make_docx(title, text),
            file_name=filename + ".docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            key=key + "_docx",
        )


def render_memory_approval(
    scene, alone, narrator
):
    key = scene["id"]

    with st.expander(
        "MEMORIA DEL PERSONAGGIO"
    ):
        st.caption(
            "La memoria non si aggiorna "
            "a ogni battuta. Solo dopo "
            "l'approvazione del racconto."
        )

        record = narrative_record(
            scene, alone, narrator
        )

        if not record.get("approved"):
            st.info(
                "Approva prima la scena."
            )
            return

        proposal_key = (
            "memory_proposal_" + key
            + "_" + narrator
        )

        if st.button(
            "ESTRAI RICORDI DALLA SCENA",
            key="extract_" + key + narrator,
            disabled=AI is None,
        ):
            with st.spinner(
                "Analisi dei ricordi..."
            ):
                proposal, error = propose_memory(
                    scene, alone, narrator
                )

            if error:
                st.error(error)
            else:
                st.session_state[
                    proposal_key
                ] = proposal

        proposal = st.session_state.get(
            proposal_key
        )

        if not isinstance(proposal, dict):
            return

        edited = st.text_area(
            "Modifica la scheda prima di approvarla",
            value=json_text(proposal),
            height=260,
            key="memory_editor_" + key + narrator,
        )

        if st.button(
            "APPROVA E AGGIORNA MEMORIA",
            key="approve_memory_" + key + narrator,
        ):
            parsed = parse_json(edited)

            if not isinstance(parsed, dict):
                st.error("JSON non valido.")

            elif approve_memory(
                narrator, parsed
            ):
                st.session_state.pop(
                    proposal_key, None
                )
                st.rerun()


def render_workshop(scene, alone, narrator):
    key = scene["id"]
    record = narrative_record(
        scene, alone, narrator
    )

    st.markdown("---")
    st.header("Officina del romanzo")

    notes = st.text_area(
        "Indicazioni letterarie",
        value=record.get("feedback", ""),
        key="writer_notes_" + key + narrator,
    )

    if st.button(
        "GENERA VERSIONI A E B",
        key="generate_" + key + narrator,
        disabled=AI is None,
    ):
        with st.spinner(
            "Scrittura delle due versioni..."
        ):
            a = write_story(
                scene, alone,
                narrator, "A", notes
            )

            b = (
                write_story(
                    scene, alone,
                    narrator, "B", notes
                )
                if not ai_failed(a)
                else ""
            )

        if ai_failed(a):
            st.error(a)

        elif ai_failed(b):
            st.error(b)

        else:
            def change():
                record["A"] = a
                record["B"] = b
                record["feedback"] = notes

            if save(change):
                st.rerun()

    if record.get("A") or record.get("B"):
        tabs = st.tabs([
            "VERSIONE A",
            "VERSIONE B",
        ])

        with tabs[0]:
            story_html(record.get("A", ""))

        with tabs[1]:
            story_html(record.get("B", ""))

        if st.button(
            "GIUDIZIO CRITICO",
            key="critic_" + key + narrator,
            disabled=AI is None,
        ):
            with st.spinner(
                "Valutazione critica..."
            ):
                review = critic(record)

            if ai_failed(review):
                st.error(review)
            elif save(
                lambda: record.update(
                    review=review
                )
            ):
                st.rerun()

    if record.get("review"):
        st.write(record["review"])

    options = ["Bozza"]

    if record.get("A"):
        options.append("Versione A")

    if record.get("B"):
        options.append("Versione B")

    selected = st.selectbox(
        "Testo da revisionare",
        options,
        key="version_" + key + narrator,
    )

    initial = (
        record.get("A", "")
        if selected == "Versione A"
        else record.get("B", "")
        if selected == "Versione B"
        else record.get("draft", "")
    )

    editor_key = (
        "editor_"
        + key + "_"
        + narrator + "_"
        + normalize(selected)
    )

    source_key = editor_key + "_source"

    if st.session_state.get(source_key) != initial:
        st.session_state[editor_key] = initial
        st.session_state[source_key] = initial

    edited = st.text_area(
        "Racconto",
        height=320,
        key=editor_key,
    )

    reason = st.text_input(
        "Motivazione",
        key="reason_" + key + narrator,
    )

    left, right = st.columns(2)

    with left:
        if st.button(
            "SALVA BOZZA",
            key="draft_" + key + narrator,
        ):
            if save(
                lambda: record.update(
                    draft=edited,
                    status="bozza",
                )
            ):
                st.success("Bozza salvata.")

    with right:
        if st.button(
            "APPROVA E MEMORIZZA",
            key="approve_" + key + narrator,
            disabled=not edited.strip(),
        ):
            def change():
                previous = record.get(
                    "approved", ""
                )

                if previous and previous != edited:
                    record["history"].append({
                        "text": previous,
                        "replaced_at": now(),
                        "reason": reason,
                    })

                record["approved"] = edited
                record["draft"] = edited
                record["approved_at"] = now()
                record["status"] = "approvata"

                log_change(
                    "testo approvato",
                    scene["title"],
                    previous,
                    edited,
                    reason,
                )

            if save(change):
                st.rerun()

    export_story(
        scene["title"],
        key + "_" + narrator.lower(),
        edited,
        "export_" + key + narrator,
    )

    if record.get("approved"):
        with st.expander("TESTO APPROVATO"):
            story_html(
                record["approved"]
            )

    with st.expander("VERSIONI PRECEDENTI"):
        for item in reversed(
            as_list(record.get("history"))
        ):
            st.caption(
                item.get("replaced_at", "")
            )
            st.code(
                text_diff(
                    item.get("text", ""),
                    record.get("approved", ""),
                ),
                language="diff",
            )

    with st.expander("TRASCRIZIONE"):
        st.text(
            scene_transcript(scene, alone)
        )

    render_memory_approval(
        scene, alone, narrator
    )


# ============================================================
# SCENE AUTONOME DI ADA
# ============================================================


def render_ada_scene_manager():
    st.sidebar.markdown("---")
    st.sidebar.subheader("SCENE ADA ALONE")

    if st.sidebar.button(
        "+ NUOVA SCENA ADA ALONE",
        key="new_ada_scene",
        use_container_width=True,
    ):
        if save(
            lambda: book["ada_scenes"].append(
                empty_ada_scene()
            )
        ):
            st.session_state.ada_scene_id = (
                book["ada_scenes"][-1]["id"]
            )
            st.rerun()

    if book["ada_scenes"]:
        identifiers = [
            s["id"] for s in book["ada_scenes"]
        ]

        current = st.session_state.ada_scene_id

        if current not in identifiers:
            current = identifiers[0]

        selected = st.sidebar.selectbox(
            "SCENA PRIVATA DI ADA",
            identifiers,
            index=identifiers.index(current),
            format_func=lambda scene_id: (
                ada_scene(scene_id)["title"]
            ),
            key="ada_scene_select",
        )

        st.session_state.ada_scene_id = selected

    else:
        st.sidebar.caption(
            "Nessuna scena autonoma creata."
        )


def render_shared_navigation():
    index = st.session_state.scene_index

    selected = st.sidebar.selectbox(
        "SCENA PRINCIPALE",
        SCENE_IDS,
        index=index,
        format_func=lambda sid: (
            f"{scene_index(sid)+1:02d}"
            " — "
            + SCENE_NAMES[sid]
        ),
        key="main_scene_selector",
    )

    if selected != SCENE_IDS[index]:
        st.session_state.scene_index = (
            scene_index(selected)
        )
        st.rerun()


# ============================================================
# PAGINE PERSONAGGI E ARCHIVIO
# ============================================================


def page_people():
    apply_theme(False)

    character = st.selectbox(
        "Personaggio",
        PEOPLE,
    )

    render_header(
        character,
        "Profilo di " + character,
    )

    agent = book["agents"][character]

    st.write(agent["profile"])

    new_profile = st.text_area(
        "Modifica profilo",
        value=agent["profile"],
        height=180,
        key="profile_" + character,
    )

    if st.button(
        "APPROVA PROFILO",
        key="approve_profile_" + character,
    ):
        def change():
            old = agent["profile"]
            agent["profile_history"].append({
                "profile": old,
                "date": now(),
            })
            agent["profile"] = new_profile

            log_change(
                "profilo",
                character,
                old,
                new_profile,
            )

        if save(change):
            st.rerun()

    st.subheader("Memoria narrativa")

    memory = agent["smart_memory"]

    identity = st.text_area(
        "Identità narrativa sintetica",
        value=memory.get("identity", ""),
        height=120,
        key="identity_" + character,
    )

    if st.button(
        "SALVA IDENTITÀ NARRATIVA",
        key="save_identity_" + character,
    ):
        if save(
            lambda: memory.update(
                identity=identity
            )
        ):
            st.rerun()

    with st.expander("RICORDI"):
        for event in as_list(
            memory.get("events")
        ):
            st.write(
                event.get("text", "")
            )

    with st.expander("RELAZIONI"):
        for name, value in as_dict(
            memory.get("relations")
        ).items():
            st.write(
                "**" + name + ":** " + value
            )

    with st.expander("SEGRETI CONOSCIUTI"):
        for item in as_list(
            memory.get("secrets")
        ):
            st.write(
                item.get("text", "")
            )


def page_archive():
    apply_theme(False)
    st.title("Archivio del romanzo")

    st.subheader("Filone principale")

    for scene in book["scenes"]:
        for person in PEOPLE:
            approved = as_dict(
                scene["narrative"].get(person)
            ).get("approved", "")

            if approved:
                with st.expander(
                    scene["title"]
                    + " — "
                    + person
                ):
                    story_html(approved)

    st.subheader("Filone esclusivo ADA ALONE")

    for scene in book["ada_scenes"]:
        approved = scene[
            "narrative"
        ].get("approved", "")

        if approved:
            with st.expander(
                scene["title"]
                + " — ADA ALONE"
            ):
                story_html(approved)


def page_editorial():
    apply_theme(False)
    st.title("Registro editoriale")

    for item in reversed(
        book["editorial_log"]
    ):
        with st.expander(
            item.get("date", "")[:10]
            + " — "
            + item.get("subject", "")
        ):
            st.write(
                item.get("category", "")
            )
            st.write(
                item.get("reason", "")
            )
            st.code(
                text_diff(
                    item.get("before", ""),
                    item.get("after", ""),
                ),
                language="diff",
            )


# ============================================================
# IMPOSTAZIONI
# ============================================================


def page_settings():
    apply_theme(False)
    st.title("Economia API e memoria")

    st.write(
        "GO ALONE mantiene quattro chiamate "
        "autonome, una per ogni intervento. "
        "Qui puoi regolare la quantità di "
        "contesto inviata agli agenti."
    )

    settings = book["settings"]

    lines = st.slider(
        "Battute recenti per richiesta",
        min_value=4,
        max_value=20,
        value=int(
            settings.get("context_lines", 10)
        ),
    )

    memories = st.slider(
        "Ricordi pertinenti per agente",
        min_value=2,
        max_value=12,
        value=int(
            settings.get("memory_items", 8)
        ),
    )

    if st.button("SALVA IMPOSTAZIONI"):
        if save(
            lambda: settings.update(
                context_lines=lines,
                memory_items=memories,
            )
        ):
            st.rerun()

    st.caption(
        "Ridurre questi valori diminuisce "
        "i token inviati, ma può togliere "
        "informazioni necessarie. "
        "Il valore predefinito è equilibrato."
    )


# ============================================================
# INGRESSO
# ============================================================


def page_entrance():
    apply_theme(False)

    image = find_asset("TuringHotel.jpg")

    if image:
        st.image(
            str(image),
            use_container_width=True,
        )

    st.title("TURING HOTEL")
    st.caption("Un romanzo a più coscienze")

    if st.button(
        "ENTRA AL TURING HOTEL",
        use_container_width=True,
    ):
        st.session_state.page = "characters"
        st.rerun()


def page_characters():
    apply_theme(False)
    st.title("Scegli una storia")

    columns = st.columns(2)

    for column, person in zip(
        columns, ["Noa", "Ada"]
    ):
        with column:
            image = portrait_path(person)

            if image:
                st.image(
                    str(image),
                    use_container_width=True,
                )

            st.subheader(person)

            if st.button(
                "ENTRA NELLA STORIA DI "
                + person.upper(),
                key="choose_" + person,
                use_container_width=True,
            ):
                st.session_state.pov = person
                st.session_state.page = "story"
                st.rerun()


# ============================================================
# BARRA LATERALE
# ============================================================


def render_sidebar():
    st.sidebar.title("TURING HOTEL")

    pov = st.session_state.pov

    if pov == "Ada":
        view = st.sidebar.radio(
            "FILONE DI ADA",
            [
                "Storia condivisa",
                "ADA ALONE",
            ],
            key="ada_view",
        )

        if view == "ADA ALONE":
            render_ada_scene_manager()
        else:
            render_shared_navigation()
    else:
        render_shared_navigation()

    st.sidebar.markdown("---")

    for label, target in [
        ("ARCHIVIO SCENE", "archive"),
        ("ARCHIVIO PERSONAGGI", "people"),
        ("REGISTRO EDITORIALE", "editorial"),
        ("ECONOMIA API", "settings"),
        ("CAMBIA STORIA", "characters"),
    ]:
        if st.sidebar.button(
            label,
            key="nav_" + target,
            use_container_width=True,
        ):
            st.session_state.page = target
            st.rerun()

    with st.sidebar.expander(
        "ARCHIVIO GITHUB"
    ):
        st.caption(
            GH_OWNER + "/" + GH_REPO
        )

        if st.button("RICARICA ARCHIVIO"):
            if reload_book():
                st.rerun()

        st.download_button(
            "SCARICA BACKUP",
            json_text(book).encode("utf-8"),
            file_name="turing_hotel_backup.json",
            mime="application/json",
        )

        pending = st.session_state.pending_backup

        if pending is not None:
            st.warning(
                "È disponibile una copia "
                "dell'ultima modifica non salvata."
            )

            st.download_button(
                "RECUPERA MODIFICA NON SALVATA",
                json_text(pending).encode("utf-8"),
                file_name="modifica_non_salvata.json",
                mime="application/json",
            )

    with st.sidebar.expander(
        "PRIMA BOZZA"
    ):
        load_reference()

        if st.session_state.reference_text:
            st.success(
                "Prima bozza collegata."
            )
        else:
            st.warning(
                st.session_state.reference_error
                or "Reference non disponibile."
            )


# ============================================================
# PAGINA STORIA
# ============================================================


def page_story():
    render_sidebar()

    pov = st.session_state.pov

    alone = (
        pov == "Ada"
        and st.session_state.get(
            "ada_view", "Storia condivisa"
        ) == "ADA ALONE"
    )

    apply_theme(alone)

    if alone:
        scene = ada_scene(
            st.session_state.ada_scene_id
        )

        if scene is None:
            st.info(
                "Per iniziare una storia autonoma "
                "premi + NUOVA SCENA ADA ALONE "
                "nella barra laterale."
            )
            return

        narrator = "Ada"

    else:
        sid = SCENE_IDS[
            st.session_state.scene_index
        ]

        scene = main_scene(sid)

        cast = scene_cast(scene)

        narrator = (
            pov if pov in cast
            else (
                cast[0] if cast
                else pov
            )
        )

    render_header(
        pov,
        scene["title"],
        alone=alone,
    )

    if alone:
        st.info(
            "Questa scena appartiene "
            "soltanto alla storia di Ada. "
            "Può avere una durata e una "
            "posizione proprie nella cronologia."
        )

    st.markdown("---")

    render_media(
        scene,
        pov,
        alone,
    )

    st.markdown("---")

    render_context(
        scene,
        alone,
    )

    render_advisor(
        scene,
        alone,
    )

    st.markdown("---")

    render_dialogues(
        scene,
        alone,
    )

    st.markdown("---")

    if not alone:
        options = scene_cast(scene)

        if options:
            narrator = st.selectbox(
                "PUNTO DI VISTA",
                options,
                index=(
                    options.index(pov)
                    if pov in options
                    else 0
                ),
                key="narrator_" + scene["id"],
            )

    render_workshop(
        scene,
        alone,
        narrator,
    )

    st.markdown("---")

    if alone:
        current_index = next(
            (
                i
                for i, item in enumerate(
                    book["ada_scenes"]
                )
                if item["id"] == scene["id"]
            ),
            0,
        )

        left, right = st.columns(2)

        with left:
            if current_index > 0:
                if st.button("← ADA PRECEDENTE"):
                    st.session_state.ada_scene_id = (
                        book["ada_scenes"][
                            current_index - 1
                        ]["id"]
                    )
                    st.rerun()

        with right:
            if current_index < len(
                book["ada_scenes"]
            ) - 1:
                if st.button("ADA SUCCESSIVA →"):
                    st.session_state.ada_scene_id = (
                        book["ada_scenes"][
                            current_index + 1
                        ]["id"]
                    )
                    st.rerun()

    else:
        index = st.session_state.scene_index
        left, right = st.columns(2)

        with left:
            if index > 0:
                if st.button("← PRECEDENTE"):
                    st.session_state.scene_index -= 1
                    st.rerun()

        with right:
            if index < 30:
                if st.button("SUCCESSIVA →"):
                    st.session_state.scene_index += 1
                    st.rerun()


# ============================================================
# ARCHIVIO MANCANTE
# ============================================================


if not st.session_state.archive_exists:
    st.title("Archivio non trovato")

    st.warning(
        "romanzo_web.json non è stato trovato. "
        "Non creare un nuovo archivio se "
        "dovrebbe già esistere."
    )

    confirmed = st.checkbox(
        "Confermo la creazione "
        "di un nuovo archivio."
    )

    if st.button(
        "CREA ARCHIVIO",
        disabled=not confirmed,
    ):
        try:
            sha = put_github_file(
                GH_OWNER,
                GH_REPO,
                GH_BRANCH,
                GH_TOKEN,
                ARCHIVE_PATH,
                json_text(book).encode("utf-8"),
            )

            st.session_state.archive_sha = sha
            st.session_state.saved_base = (
                copy.deepcopy(book)
            )
            st.session_state.archive_exists = True
            st.rerun()

        except Exception as exc:
            st.error(str(exc))

    st.stop()


# ============================================================
# NAVIGAZIONE PRINCIPALE
# ============================================================


if st.session_state.page in (
    "people",
    "archive",
    "editorial",
    "settings",
):
    if st.button("← TORNA ALLA STORIA"):
        st.session_state.page = "story"
        st.rerun()


if st.session_state.page == "entrance":
    page_entrance()

elif st.session_state.page == "characters":
    page_characters()

elif st.session_state.page == "story":
    page_story()

elif st.session_state.page == "people":
    page_people()

elif st.session_state.page == "archive":
    page_archive()

elif st.session_state.page == "editorial":
    page_editorial()

elif st.session_state.page == "settings":
    page_settings()

else:
    st.session_state.page = "entrance"
    st.rerun()
